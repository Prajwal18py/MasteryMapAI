import io, math, os
from fastapi import HTTPException


def ocr_image(raw):
    if os.getenv("CLOUD_PROFILE") == "lite":
        raise HTTPException(503, "Image/scanned-PDF OCR is disabled in the lite hosting profile. Upload a text PDF, DOCX, PPTX or TXT, or use the full local app.")
    try:
        import numpy as np
        from PIL import Image
        from rapidocr_onnxruntime import RapidOCR

        image = Image.open(io.BytesIO(raw))
        image.thumbnail((2000, 2000))
        engine = RapidOCR()
        result, _ = engine(np.array(image.convert("RGB")))
        return "\n".join(row[1] for row in result or [])
    except ImportError:
        raise HTTPException(503, "OCR dependencies are missing. Run setup again.")
    except Exception:
        raise HTTPException(
            400, "Could not read text from this image. Use a clear, upright image."
        )


def extract(name, raw):
    name = name.lower()
    pages = None
    method = "text extraction"
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(raw))
            pages = len(reader.pages)
            if pages > 100:
                raise HTTPException(400, "Split PDFs into at most 100 pages.")
            text = "\n\n".join(
                f"[Page {i+1}]\n" + (p.extract_text() or "")
                for i, p in enumerate(reader.pages)
            )
            if len("".join(p.extract_text() or "" for p in reader.pages).strip()) < 30:
                if pages > 12:
                    raise HTTPException(
                        400, "For scanned PDFs, upload at most 12 pages at a time."
                    )
                if os.getenv("CLOUD_PROFILE") == "lite":
                    raise HTTPException(503, "Scanned-PDF OCR is disabled in the lite hosting profile. Upload a text-based document or use the full local app.")
                import fitz

                doc = fitz.open(stream=raw, filetype="pdf")
                text = "\n\n".join(
                    f"[Page {i+1}]\n"
                    + ocr_image(
                        p.get_pixmap(matrix=fitz.Matrix(1.7, 1.7)).tobytes("png")
                    )
                    for i, p in enumerate(doc)
                )
                method = "OCR"
        elif name.endswith(".docx"):
            from docx import Document

            d = Document(io.BytesIO(raw))
            text = (
                "\n".join(p.text for p in d.paragraphs)
                + "\n"
                + "\n".join(
                    " | ".join(c.text for c in r.cells)
                    for t in d.tables
                    for r in t.rows
                )
            )
        elif name.endswith(".pptx"):
            from pptx import Presentation

            d = Presentation(io.BytesIO(raw))
            pages = len(d.slides)
            text = "\n\n".join(
                f"[Slide {i+1}]\n"
                + "\n".join(s.text for s in slide.shapes if s.has_text_frame)
                for i, slide in enumerate(d.slides)
            )
        elif name.endswith((".png", ".jpg", ".jpeg", ".webp")):
            text = ocr_image(raw)
            method = "OCR"
        elif name.endswith((".txt", ".md")):
            text = raw.decode("utf-8-sig")
        else:
            raise HTTPException(400, "Use PDF, DOCX, PPTX, TXT, Markdown, PNG or JPEG.")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            400,
            "Could not extract this document. Check that it is not encrypted or corrupt.",
        )
    if not 20 <= len(text.strip()) <= 250000:
        raise HTTPException(
            400, "Document must contain 20–250,000 extracted characters."
        )
    return text, {
        "pages": pages,
        "method": method,
        "passages": math.ceil(len(text) / 850),
        "status": "Ready",
    }
