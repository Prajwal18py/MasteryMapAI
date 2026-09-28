import os, tempfile, unittest, uuid, json
from datetime import datetime, timezone, timedelta
from unittest.mock import patch, AsyncMock

TEMP = tempfile.TemporaryDirectory()
os.environ["MASTERYMAP_DATA"] = TEMP.name
os.environ["ENABLE_CODE_RUNNER"] = "false"
from fastapi.testclient import TestClient
from app.main import app
from app.database import db


class StudentTests(unittest.TestCase):
    def setUp(self):
        self.c = TestClient(app)
        self.c.__enter__()
        r = self.c.post(
            "/api/auth/register",
            json={
                "name": "Learner",
                "email": uuid.uuid4().hex + "@example.test",
                "password": "testing-password",
            },
        )
        self.assertEqual(r.status_code, 200, r.text)
        self.sid = self.c.get("/api/subjects").json()[0]["id"]
        self.b = "/api/subjects/" + self.sid

    def tearDown(self):
        self.c.__exit__(None, None, None)

    def test_practice_idempotency(self):
        q = self.c.post(self.b + "/practice", json={}).json()
        self.assertNotIn("correct", q)
        p = {"questionId": q["id"], "choice": 0, "attemptId": uuid.uuid4().hex}
        self.assertEqual(self.c.post(self.b + "/answer", json=p).status_code, 200)
        self.c.post(self.b + "/answer", json=p)
        self.assertEqual(self.c.get(self.b).json()["learning"]["answers"], 1)

    def test_privacy(self):
        with TestClient(app) as other:
            other.post(
                "/api/auth/register",
                json={
                    "name": "Other",
                    "email": uuid.uuid4().hex + "@example.test",
                    "password": "testing-password",
                },
            )
            self.assertEqual(other.get(self.b).status_code, 404)

    def test_custom_subject(self):
        r = self.c.post(
            "/api/subjects",
            json={
                "name": "Algebra",
                "concepts": [
                    {
                        "id": "vectors",
                        "name": "Vectors",
                        "summary": "Vectors are ordered lists of coordinates.",
                    }
                ],
            },
        )
        self.assertEqual(r.status_code, 200, r.text)
        b = "/api/subjects/" + r.json()["id"]
        self.assertIsNone(self.c.get(b).json()["learning"]["overall"])
        self.assertFalse(self.c.get(b + "/models").json()["available"])

    def test_cycles(self):
        r = self.c.post(
            "/api/subjects",
            json={
                "name": "Cyclic",
                "concepts": [
                    {"id": "aa", "name": "AA", "prereqs": ["bb"]},
                    {"id": "bb", "name": "BB", "prereqs": ["aa"]},
                ],
            },
        )
        self.assertEqual(r.status_code, 422)

    def test_material_chat(self):
        r = self.c.post(
            self.b + "/materials",
            files={
                "file": (
                    "notes.txt",
                    b"Functions accept arguments and return results. Closures capture variables.",
                    "text/plain",
                )
            },
        )
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(len(self.c.get(self.b + "/materials").json()), 1)
        r2 = self.c.post(
            self.b + "/tutor",
            json={"message": "Explain functions", "referenceOnly": True},
        )
        self.assertEqual(r2.status_code, 200, r2.text)
        self.assertEqual(len(self.c.get(self.b + "/chat").json()), 2)
        self.assertEqual(
            self.c.delete("/api/materials/" + r.json()["id"]).status_code, 200
        )

    def test_quiz_deadline(self):
        r = self.c.post(
            self.b + "/quizzes", json={"title": "Test quiz", "count": 2, "duration": 1}
        )
        self.assertEqual(r.status_code, 200, r.text)
        b = "/api/quizzes/" + r.json()["id"]
        q = self.c.post(b + "/start").json()
        qid = q["questions"][0]["id"]
        self.c.post(b + "/submit", json={"answers": {qid: 0}})
        with db() as c:
            c.execute(
                "UPDATE quizzes SET started_at=? WHERE id=?",
                (
                    (datetime.now(timezone.utc) - timedelta(minutes=2)).isoformat(),
                    q["id"],
                ),
            )
        result = self.c.post(
            b + "/submit", json={"answers": {qid: 1}, "finish": True}
        ).json()["result"]
        self.assertTrue(result["expired"])
        self.assertEqual(result["feedback"][0]["choice"], 0)
        self.c.post(b + "/submit", json={"finish": True})
        self.assertEqual(self.c.get(self.b).json()["learning"]["answers"], 1)

    def test_rule_plan_approval(self):
        r = self.c.post(self.b + "/runs", json={"minutes": 37})
        self.assertEqual(r.status_code, 200, r.text)
        trace = self.c.get(self.b + "/runs").json()[0]
        self.assertEqual(trace["status"], "awaiting_approval", trace)
        self.assertIsNone(self.c.get(self.b + "/plan").json())
        p = self.c.post("/api/runs/" + trace["id"] + "/approve").json()
        self.assertEqual(sum(x["minutes"] for x in p["items"]), 37)

    def test_critic_revision(self):
        ids = [x["id"] for x in self.c.get(self.b).json()["subject"]["concepts"]]
        with patch(
            "app.student_agents.generate",
            new=AsyncMock(
                side_effect=[
                    json.dumps({"concepts": [ids[-1]], "rationale": "Review"}),
                    json.dumps(
                        {"approved": False, "reason": "Repair", "replacement": [ids[0]]}
                    ),
                ]
            ),
        ):
            self.c.post(self.b + "/runs", json={"useAI": True})
        trace = self.c.get(self.b + "/runs").json()[0]
        self.assertEqual(trace["status"], "awaiting_approval", trace)
        self.assertTrue(
            any(e["action"] == "Apply critic revision" for e in trace["events"])
        )

    def test_ai_proposal_not_saved(self):
        s = self.c.get(self.b).json()["subject"]
        q = s["questions"][0]
        with patch("app.main.generate", new=AsyncMock(return_value=json.dumps([q]))):
            r = self.c.post(
                self.b + "/questions/generate",
                json={"concept": q["concept"], "count": 1},
            )
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(
            len(self.c.get(self.b).json()["subject"]["questions"]), len(s["questions"])
        )

    def test_origin_and_code(self):
        self.assertEqual(
            self.c.post(
                self.b + "/practice",
                json={},
                headers={"origin": "https://untrusted.test"},
            ).status_code,
            403,
        )
        self.assertEqual(
            self.c.post(
                "/api/code/run",
                json={"exercise": "sum", "code": "def solve(x): return sum(x)"},
            ).status_code,
            410,
        )

    def test_password_export(self):
        self.assertEqual(self.c.get("/api/export").status_code, 200)
        self.assertEqual(
            self.c.post(
                "/api/auth/password",
                json={
                    "current": "testing-password",
                    "replacement": "new-testing-password",
                },
            ).status_code,
            200,
        )
        self.assertEqual(self.c.get("/api/auth/me").status_code, 200)

    def test_document_formats(self):
        import io, fitz
        from docx import Document
        from pptx import Presentation
        from app.extraction import extract

        d = Document()
        d.add_paragraph("Functions accept arguments and return values.")
        b = io.BytesIO()
        d.save(b)
        self.assertIn("arguments", extract("a.docx", b.getvalue())[0])
        p = Presentation()
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = "Functions accept arguments and return values."
        b = io.BytesIO()
        p.save(b)
        self.assertIn("arguments", extract("a.pptx", b.getvalue())[0])
        p = fitz.open()
        s = p.new_page()
        s.insert_text((30, 30), "Functions accept arguments and return values.")
        self.assertIn("arguments", extract("a.pdf", p.tobytes())[0])
